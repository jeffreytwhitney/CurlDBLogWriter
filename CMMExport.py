import os

import openpyxl

from IExport import IExport
from openpyxl import Workbook


class CMMExport(IExport):
    _file_path: str
    _job_number: str
    _part_number: str
    _sequence_numbers: list[int] = []
    _file_lines: list[str]

    def __init__(self, file_path: str):
        super().__init__(file_path)

        self._file_path = file_path
        xl_file = openpyxl.open(file_path)
        ws = xl_file.active
        self._part_number = ws['B2'].value
        self._job_number = ws['C11'].value
        self._sequence_numbers.append(int(ws['C12'].value))

    @property
    def file_path(self) -> str:
        return self._file_path

    @property
    def job_number(self) -> str:
        return self._job_number

    @property
    def part_number(self) -> str:
        return self._part_number

    @property
    def sequence_numbers(self) -> list[int]:
        return self._sequence_numbers

    @property
    def file_name(self) -> str:
        return os.path.basename(self._file_path)
